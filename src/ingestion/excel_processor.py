import csv
from pathlib import Path

from openpyxl import load_workbook


# Rows are grouped so each chunk carries enough context but stays
# small enough to rank precisely.
MAX_ROWS_PER_GROUP = 25

MAX_CHARS_PER_GROUP = 1500


def _format_value(value) -> str:

    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def rows_to_documents(rows, file_path, file_type, sheet=None) -> list[dict]:
    """
    First row is the header. Each output document holds a group of rows
    written as 'Row N: Header: value | Header: value'.
    """

    rows = [
        [_format_value(value) for value in row]
        for row in rows
    ]

    rows = [row for row in rows if any(row)]

    if not rows:
        return []

    header = [
        name or f"Column {index + 1}"
        for index, name in enumerate(rows[0])
    ]

    prefix = f"Sheet: {sheet}\n" if sheet else ""
    columns = "Columns: " + ", ".join(header)

    documents = []
    lines = []
    group = {"start": 2, "end": 2}

    def flush():
        if not lines:
            return

        documents.append(
            {
                "text": prefix + columns + "\n" + "\n".join(lines),
                "metadata": {
                    "source": file_path,
                    "file_type": file_type,
                    "sheet": sheet,
                    "row_start": group["start"],
                    "row_end": group["end"],
                    # Already sized; the chunker must not split rows apart.
                    "presplit": True,
                },
            }
        )

        lines.clear()

    for row_number, row in enumerate(rows[1:], start=2):

        pairs = [
            f"{name}: {value}"
            for name, value in zip(header, row)
            if value
        ]

        line = f"Row {row_number}: " + " | ".join(pairs)

        too_big = (
            len(lines) >= MAX_ROWS_PER_GROUP
            or sum(len(text) for text in lines) + len(line) > MAX_CHARS_PER_GROUP
        )

        if lines and too_big:
            flush()
            group["start"] = row_number

        lines.append(line)
        group["end"] = row_number

    flush()

    # A sheet with only a header row is still searchable by column names.
    if not documents:
        documents.append(
            {
                "text": prefix + columns,
                "metadata": {
                    "source": file_path,
                    "file_type": file_type,
                    "sheet": sheet,
                    "presplit": True,
                },
            }
        )

    return documents


def extract_excel_content(file_path: str, progress=None) -> list[dict]:

    workbook = load_workbook(file_path, read_only=True, data_only=True)

    documents = []

    try:
        for worksheet in workbook.worksheets:

            if progress:
                progress(f"Reading sheet {worksheet.title}")

            documents.extend(
                rows_to_documents(
                    worksheet.iter_rows(values_only=True),
                    file_path,
                    "excel",
                    sheet=worksheet.title,
                )
            )
    finally:
        workbook.close()

    return documents


def extract_csv_content(file_path: str, progress=None) -> list[dict]:

    with open(
        file_path,
        newline="",
        encoding="utf-8-sig",
        errors="replace",
    ) as file:

        sample = file.read(4096)
        file.seek(0)

        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel

        rows = list(csv.reader(file, dialect))

    return rows_to_documents(
        rows,
        file_path,
        "csv",
        sheet=Path(file_path).stem,
    )
