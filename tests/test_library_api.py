import io
import json

import docx
import openpyxl

from backend import database, library


def upload(client, headers, name, content, mime="application/octet-stream"):
    return client.post(
        "/api/sources",
        files=[("files", (name, content, mime))],
        headers=headers,
    )


def sources(client, headers):
    return client.get("/api/sources", headers=headers).json()["sources"]


def chat_with(client, headers, question, **body):
    """(conversation, events after the conversation event)."""
    response = client.post("/api/chat/stream", json={"question": question, **body}, headers=headers)
    assert response.status_code == 200, response.text
    events = [json.loads(line) for line in response.text.splitlines() if line.strip()]
    assert events[0]["type"] == "conversation"
    return events[0]["conversation"], events[1:]


def chat(client, headers, question, **body):
    return chat_with(client, headers, question, **body)[1]


def messages(client, headers, conversation_id):
    return client.get(f"/api/conversations/{conversation_id}", headers=headers).json()["messages"]


def xlsx_bytes():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Q1"
    sheet.append(["Region", "Revenue", "Units"])
    sheet.append(["North", 120000, 340])
    sheet.append(["South", 98000, 290])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def docx_bytes():
    document = docx.Document()
    document.add_heading("Leave policy", level=1)
    document.add_paragraph("Employees receive 24 days of paid annual leave.")
    document.add_heading("Remote work", level=1)
    document.add_paragraph("Remote work is allowed two days per week.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "Days"
    table.rows[1].cells[0].text = "Engineer"
    table.rows[1].cells[1].text = "3"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------- ingestion ----------------

def test_upload_multiple_types_become_ready(client, alice):

    files = [
        ("notes.txt", b"Project Falcon launches on 14 March 2027.", "text/plain"),
        ("sales.xlsx", xlsx_bytes(), "application/vnd.ms-excel"),
        ("policy.docx", docx_bytes(), "application/msword"),
        ("people.csv", b"name,team\nPriya,Mobile\nArun,Web\n", "text/csv"),
        ("page.html", b"<html><title>Guide</title><body><h1>Setup</h1><p>Install the agent first.</p>"
                      b"<table><tr><th>OS</th><th>Steps</th></tr><tr><td>Linux</td><td>apt install</td></tr></table>"
                      b"</body></html>", "text/html"),
    ]

    for name, content, mime in files:
        assert upload(client, alice, name, content, mime).status_code == 201

    listed = {source["name"]: source for source in sources(client, alice)}

    assert set(listed) == {name for name, _, _ in files}
    for source in listed.values():
        assert source["status"] == "ready", source
        assert source["chunk_count"] >= 1

    assert listed["sales.xlsx"]["file_type"] == "excel"
    assert listed["page.html"]["chunk_count"] >= 2  # prose + table rows


def test_answers_cite_the_right_source_and_location(client, alice):

    upload(client, alice, "sales.xlsx", xlsx_bytes())
    upload(client, alice, "notes.txt", b"The office cafeteria opens at 8am.")

    events = chat(client, alice, "What was the revenue for the South region?")
    kinds = [event["type"] for event in events]

    assert kinds[0] == "sources" and kinds[-1] == "done" and "token" in kinds

    top = events[0]["sources"][0]
    assert top["name"] == "sales.xlsx"
    assert top["location"] == "Q1 · rows 2–3"
    assert "South" in top["snippet"]
    assert "text" not in top  # full chunk text is not sent to the browser
    assert "98000" in client.calls["prompts"][-1]


def test_unsupported_and_empty_requests_rejected(client, alice):
    assert upload(client, alice, "virus.exe", b"MZ").status_code == 400
    assert client.post("/api/sources", headers=alice).status_code == 400
    bad_url = client.post("/api/sources", data={"urls": "ftp://x"}, headers=alice)
    assert bad_url.status_code == 400


def test_upload_size_limit(client, alice, monkeypatch):
    from config import settings
    monkeypatch.setattr(settings, "MAX_UPLOAD_MB", 1)
    response = upload(client, alice, "big.txt", b"a" * (1024 * 1024 + 10))
    assert response.status_code == 413
    assert sources(client, alice) == []


def test_private_urls_are_blocked(client, alice):
    client.post("/api/sources", data={"urls": "http://127.0.0.1:8000/secret"}, headers=alice)
    source = sources(client, alice)[0]
    assert source["status"] == "failed"
    assert "private or local" in source["error"]


def test_empty_file_fails_with_message_and_can_retry(client, alice, tmp_path):
    upload(client, alice, "empty.txt", b"   ")
    source = sources(client, alice)[0]
    assert source["status"] == "failed"
    assert "No text" in source["error"]

    # Fix the file on disk and retry.
    path = database.get_source(1, source["id"])["location"]
    open(path, "w").write("Now it has content about zebras.")
    retried = client.post(f"/api/sources/{source['id']}/retry", headers=alice)
    assert retried.status_code == 200
    assert sources(client, alice)[0]["status"] == "ready"


# ---------------- library management ----------------

def test_library_is_additive_and_delete_removes_from_search(client, alice):

    upload(client, alice, "zebra.txt", b"A group of zebras is called a dazzle.")
    upload(client, alice, "falcon.txt", b"Project Falcon budget is 2.5 million dollars.")

    assert len(sources(client, alice)) == 2

    events = chat(client, alice, "What is a group of zebras called?")
    assert events[0]["sources"][0]["name"] == "zebra.txt"

    zebra = next(s for s in sources(client, alice) if s["name"] == "zebra.txt")
    assert client.delete(f"/api/sources/{zebra['id']}", headers=alice).status_code == 200

    names = [s["name"] for s in sources(client, alice)]
    assert names == ["falcon.txt"]

    events = chat(client, alice, "What is a group of zebras called?")
    assert all(s["name"] != "zebra.txt" for s in events[0]["sources"])
    assert not (library.UPLOAD_ROOT / "1" / zebra["id"]).exists()


def test_deleting_last_source_removes_index(client, alice):
    upload(client, alice, "a.txt", b"alpha content")
    source_id = sources(client, alice)[0]["id"]
    client.delete(f"/api/sources/{source_id}", headers=alice)
    assert not (library.VECTORSTORE_ROOT / "1").exists()
    events = chat(client, alice, "alpha?")
    assert "no processed sources" in events[-1]["answer"]


def test_library_is_per_user(client, alice, bobby):

    upload(client, alice, "secret.txt", b"The launch code is pineapple.")
    alice_source = sources(client, alice)[0]

    assert sources(client, bobby) == []

    events = chat(client, bobby, "What is the launch code?")
    assert events[0]["sources"] == []
    assert "pineapple" not in json.dumps(events)

    # Bobby cannot touch Alice's source.
    assert client.delete(f"/api/sources/{alice_source['id']}", headers=bobby).status_code == 404
    assert client.post(f"/api/sources/{alice_source['id']}/retry", headers=bobby).status_code == 404
    assert client.get(f"/api/sources/{alice_source['id']}/link", headers=bobby).status_code == 404


def test_file_links_are_scoped_and_signed(client, alice, bobby):

    upload(client, alice, "doc.txt", b"hello world")
    source_id = sources(client, alice)[0]["id"]

    url = client.get(f"/api/sources/{source_id}/link", headers=alice).json()["url"]
    response = client.get(url)
    assert response.status_code == 200 and response.content == b"hello world"

    assert client.get(f"/api/sources/{source_id}/file?access=bad").status_code == 401

    upload(client, bobby, "other.txt", b"bobby data")
    bobby_id = sources(client, bobby)[0]["id"]
    token = url.split("access=")[1]
    assert client.get(f"/api/sources/{bobby_id}/file?access={token}").status_code == 401


def test_server_paths_are_not_exposed(client, alice):
    upload(client, alice, "notes.txt", b"Kettle descaling happens every Monday.")
    events = chat(client, alice, "When is the kettle descaled?")
    assert "uploads" not in json.dumps(events)


# ---------------- chat ----------------

def test_chat_without_sources_explains_what_to_do(client, alice):
    events = chat(client, alice, "anything")
    assert "no processed sources" in events[-1]["answer"]


def test_history_and_follow_up_rewrite(client, alice):

    upload(client, alice, "falcon.txt", b"Project Falcon is led by Dr. Meera Iyer.")

    conversation, _ = chat_with(client, alice, "Who leads Project Falcon?")
    assert client.calls["rewrites"] == []  # first question: nothing to resolve

    events = chat(client, alice, "When did she join?", conversation_id=conversation["id"])
    assert len(client.calls["rewrites"]) == 1
    assert "Who leads Project Falcon?" in client.calls["rewrites"][0]
    assert events[0]["query"] == "rewritten standalone query"

    history = messages(client, alice, conversation["id"])
    assert [m["role"] for m in history] == ["user", "assistant", "user", "assistant"]
    assert history[1]["sources"][0]["name"] == "falcon.txt"

    # A new conversation starts without that context.
    client.calls["rewrites"].clear()
    chat(client, alice, "When did she join?")
    assert client.calls["rewrites"] == []


def test_non_streaming_chat(client, alice):
    upload(client, alice, "falcon.txt", b"Project Falcon is led by Dr. Meera Iyer.")
    response = client.post("/api/chat", json={"question": "Who leads Falcon?"}, headers=alice)
    assert response.status_code == 200
    assert "Meera" in response.json()["answer"]
    assert response.json()["sources"][0]["name"] == "falcon.txt"


def test_llm_errors_are_reported_not_crashed(client, alice, monkeypatch):

    from src import rag_pipeline

    class Broken:
        def invoke(self, messages):
            raise RuntimeError("model offline")

        def stream(self, messages):
            raise RuntimeError("model offline")

    upload(client, alice, "a.txt", b"alpha beta gamma")
    monkeypatch.setattr(rag_pipeline, "get_llm", lambda config: Broken())

    conversation, events = chat_with(client, alice, "alpha beta?")
    assert events[-1] == {"type": "error", "message": "model offline"}
    assert messages(client, alice, conversation["id"]) == []


def test_restart_resumes_queued_sources(client, alice, monkeypatch):

    from backend import jobs

    submitted = []
    monkeypatch.setattr(jobs, "submit", lambda user_id, source_id: submitted.append(source_id))
    monkeypatch.setattr("backend.main.jobs.submit", jobs.submit)

    upload(client, alice, "queued.txt", b"waiting for a worker")
    source = sources(client, alice)[0]
    assert source["status"] == "queued"

    submitted.clear()
    jobs.resume_unfinished()
    assert submitted == [source["id"]]


def test_abstained_answers_have_no_sources(client, alice, monkeypatch):

    from langchain_core.messages import AIMessageChunk
    from src import rag_pipeline
    from src.generation.prompts import ABSTAIN

    class Abstains:
        def invoke(self, messages):
            return AIMessageChunk(content="query")

        def stream(self, messages):
            yield AIMessageChunk(content=ABSTAIN)

    upload(client, alice, "a.txt", b"alpha beta gamma")
    monkeypatch.setattr(rag_pipeline, "get_llm", lambda config: Abstains())

    conversation, events = chat_with(client, alice, "alpha beta?")
    assert events[0]["sources"]  # retrieval found something...
    assert events[-1]["sources"] == []  # ...but the model judged it irrelevant
    assert messages(client, alice, conversation["id"])[-1]["sources"] == []
